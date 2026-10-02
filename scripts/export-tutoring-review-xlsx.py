from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


input_path = Path(sys.argv[1])
output_path = Path(sys.argv[2])
preserve_path = Path(sys.argv[3]) if len(sys.argv) > 3 else None

with input_path.open(encoding="utf-8") as f:
    source = json.load(f)


def flatten_values(value: Any):
    if isinstance(value, dict):
        for child in value.values():
            yield from flatten_values(child)
    elif isinstance(value, list):
        for child in value:
            yield from flatten_values(child)
    else:
        yield value


def date_keys(value: Any) -> list[str]:
    text = json.dumps(value, ensure_ascii=False)
    return sorted(set(re.findall(r"20\d{2}-\d{2}-\d{2}", text)))


def leaf_count(value: Any) -> int:
    return sum(1 for item in flatten_values(value) if item not in (None, ""))


def normalise_source(source: dict) -> dict:
    # Accept either the migration template format or the original backup JSON.
    if "classInstances" in source and "occurrences" in source:
        return source

    classes = source.get("classIndex", [])
    records = source.get("records", {})
    class_instances = []
    occurrences = []
    for cls in classes:
        cid = str(cls.get("id", ""))
        rec = records.get(cid, {})
        dates = date_keys(rec.get("attendance", {}))
        class_instances.append({
            "legacyClassId": cid,
            "name": cls.get("name", ""),
            "grade": cls.get("grade", ""),
            "studentCount": len(cls.get("students", [])),
            "attendanceDateStart": dates[0] if dates else "",
            "attendanceDateEnd": dates[-1] if dates else "",
            "attendanceCellCount": leaf_count(rec.get("attendance", {})),
            "quizCellCount": leaf_count(rec.get("quiz", {})),
            "examCellCount": leaf_count(rec.get("exam", {})),
        })
        for index, student in enumerate(cls.get("students", [])):
            occurrences.append({
                "id": f"class:{cid}:roster:{index}:legacy:{student.get('id', '')}",
                "legacyStudentId": student.get("id", ""),
                "name": student.get("name", ""),
                "school": student.get("school", ""),
                "grade": cls.get("grade", ""),
                "legacyClassId": cid,
                "className": cls.get("name", ""),
                "group": student.get("group", ""),
            })
    return {"classInstances": class_instances, "occurrences": occurrences}


data = normalise_source(source)

# Preserve values from the old workbook when a user has already started filling it.
preserved_class = {}
preserved_student = {}
if preserve_path and preserve_path.exists():
    old = load_workbook(preserve_path, data_only=False)
    if "班級確認" in old.sheetnames:
        ws = old["班級確認"]
        headers = {str(c.value): c.column for c in ws[1]}
        key_col = headers.get("原始班級名稱") or headers.get("舊班級名稱")
        for row in ws.iter_rows(min_row=2):
            if key_col and row[key_col - 1].value:
                key = str(row[key_col - 1].value)
                preserved_class[key] = [c.value for c in row]
    if "學生確認" in old.sheetnames:
        ws = old["學生確認"]
        headers = {str(c.value): c.column for c in ws[1]}
        name_col = headers.get("原始姓名") or headers.get("姓名")
        class_col = headers.get("原始班級") or headers.get("班級")
        for row in ws.iter_rows(min_row=2):
            if name_col and class_col and row[name_col - 1].value:
                key = (str(row[name_col - 1].value), str(row[class_col - 1].value or ""))
                preserved_student[key] = [c.value for c in row]

wb = Workbook()
readme = wb.active
readme.title = "先看這裡"
readme_rows = [
    ["學生與班級確認表（新版）"],
    ["本版變更", "兩個同名班級不再只顯示「國中自然」，會同時顯示原始班級 ID、年級、人數與資料日期範圍。"],
    ["您不需要輸入", "不要輸入 stu_、cls_、enr_、UUID 或延續群組；正式 ID 仍由系統自動產生。"],
    ["班級確認", "請依「原始班級 ID／年級／人數／日期範圍」辨認班級，再填年度、顯示名稱與處理方式。相同名稱不代表同一班。"],
    ["班級延續", "整班升級且是同一班，選「沿用同一班」；新招生或跨學段的新班，選「建立新班」。"],
    ["學生確認", "每列是一筆舊資料中的班級成員。請依姓名、原始學校、年級、班級辨識碼判斷是否為同一位學生。"],
    ["轉學校案例", "例如王侑謙國一安南、國二德光：仍是同一個學生 ID；學校只記在該年度／班級資料，轉學不會建立新學生。"],
    ["分組", "高中數學若有分組，只填數A或數B；沒有分組請留白。"],
    ["安全規則", "任何一列仍為「待確認」都不會匯入新 Firebase；本檔只供人工確認，不會連線或修改資料庫。"],
]
for row in readme_rows:
    readme.append(row)
readme.column_dimensions["A"].width = 22
readme.column_dimensions["B"].width = 115
for row in readme.iter_rows():
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
readme.freeze_panes = "A2"

blue = PatternFill("solid", fgColor="2563EB")
yellow = PatternFill("solid", fgColor="FFF4BF")
red = PatternFill("solid", fgColor="FDE2E2")
green = PatternFill("solid", fgColor="DCFCE7")
white_bold = Font(color="FFFFFF", bold=True)


def style_header(ws):
    for cell in ws[1]:
        cell.fill = blue
        cell.font = white_bold
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


classes = wb.create_sheet("班級確認")
class_headers = [
    "原始班級 ID（辨識用）", "原始班級名稱", "原始年級", "名單人數",
    "資料日期範圍", "出缺勤格數", "平時考格數", "段考格數", "年度（請填）",
    "畫面顯示名稱（請填）", "班級處理方式", "備註",
]
classes.append(class_headers)
class_dv = DataValidation(type="list", formula1='"沿用同一班,建立新班,待確認"', allow_blank=False)
classes.add_data_validation(class_dv)
for item in data["classInstances"]:
    cid = str(item.get("legacyClassId", item.get("id", "")))
    name = str(item.get("name", item.get("className", "")))
    old = preserved_class.get(name, [])
    old_year = old[1] if len(old) > 1 else ""
    old_display = old[2] if len(old) > 2 else name
    old_action = old[3] if len(old) > 3 else "待確認"
    old_note = old[4] if len(old) > 4 else ""
    date_range = ""
    if item.get("attendanceDateStart") or item.get("attendanceDateEnd"):
        date_range = f"{item.get('attendanceDateStart', '')}～{item.get('attendanceDateEnd', '')}"
    classes.append([
        cid, name, item.get("grade", ""), item.get("studentCount", ""), date_range,
        item.get("attendanceCellCount", ""), item.get("quizCellCount", ""), item.get("examCellCount", ""),
        old_year, old_display, old_action, old_note,
    ])
style_header(classes)
for r in range(2, classes.max_row + 1):
    class_dv.add(classes.cell(r, 11))
    for c in [9, 10, 11, 12]:
        classes.cell(r, c).fill = yellow
    classes.cell(r, 11).fill = red
classes.conditional_formatting.add(f"K2:K{classes.max_row}", FormulaRule(formula=['K2="待確認"'], fill=red))
for col, width in {"A":32,"B":18,"C":12,"D":11,"E":25,"F":13,"G":13,"H":13,"I":14,"J":26,"K":18,"L":42}.items():
    classes.column_dimensions[col].width = width

students = wb.create_sheet("學生確認")
student_headers = [
    "Occurrence 顯示鍵", "舊學生 ID（辨識用）", "原始姓名", "原始學校", "原始年級",
    "原始班級 ID", "原始班級名稱", "分組", "年度（請填）", "學生處理方式",
    "既有學生目前姓名（相同才填）", "備註",
]
students.append(student_headers)
student_dv = DataValidation(type="list", formula1='"建立新學生,與既有學生相同,待確認"', allow_blank=False)
students.add_data_validation(student_dv)
for item in data["occurrences"]:
    name = str(item.get("name", ""))
    cid = str(item.get("legacyClassId", ""))
    old = preserved_student.get((name, cid), [])
    old_year = old[4] if len(old) > 4 else ""
    old_action = old[6] if len(old) > 6 else "待確認"
    old_match = old[7] if len(old) > 7 else ""
    old_note = old[8] if len(old) > 8 else ""
    students.append([
        item.get("id", ""), item.get("legacyStudentId", ""), name, item.get("school", ""),
        item.get("grade", ""), cid, item.get("className", ""), item.get("group", ""),
        old_year, old_action, old_match, old_note,
    ])
style_header(students)
for r in range(2, students.max_row + 1):
    student_dv.add(students.cell(r, 10))
    for c in [9, 10, 11, 12]:
        students.cell(r, c).fill = yellow
    students.cell(r, 10).fill = red
students.conditional_formatting.add(f"J2:J{students.max_row}", FormulaRule(formula=['J2="待確認"'], fill=red))
for col, width in {"A":48,"B":20,"C":16,"D":14,"E":12,"F":31,"G":18,"H":10,"I":14,"J":18,"K":30,"L":40}.items():
    students.column_dimensions[col].width = width

candidates = wb.create_sheet("既有學生參考")
candidates.append(["目前備份可見姓名", "提醒"])
names = sorted({str(v.get("name", "")).strip() for v in data["occurrences"] if str(v.get("name", "")).strip()})
for name in names:
    candidates.append([name, "姓名只是參考；若確認為同一人，請在學生確認頁選擇與既有學生相同。學校改變不會建立新的學生 ID。"])
style_header(candidates)
candidates.column_dimensions["A"].width = 24
candidates.column_dimensions["B"].width = 95
for row in candidates.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)

output_path.parent.mkdir(parents=True, exist_ok=True)
wb.save(output_path)
print(output_path)
