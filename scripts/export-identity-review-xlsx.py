from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import json
import sys

input_path = sys.argv[1]
output_path = sys.argv[2]
with open(input_path, encoding="utf-8") as f:
    data = json.load(f)

wb = Workbook()
ws = wb.active
ws.title = "使用說明"
notes = [
    ["學生身分確認表（系統產生 ID）"],
    ["目的", "確認哪些舊資料 occurrence 是同一位長期學生，以及哪些班級是同一條延續線。"],
    ["您不需要做的事", "不要輸入 stu_、cls_、lin_ 或 enr_；這些正式 ID 會由系統在確認後自動產生。"],
    ["學生確認方式", "在「學生確認」工作表的「確認群組」填入容易辨識的標籤，例如 C001。相同標籤代表同一位學生；不同學生使用不同標籤。這只是人工確認用標籤，不是正式 ID。"],
    ["新生", "確認群組填入新的標籤，例如 N001；系統會自動產生新的 stu_ ID。"],
    ["班級延續", "在「班級確認」的「延續判定」選「延續既有班級線」或「建立新班級線」，並在「延續群組」填入簡短標籤，例如 JPHYS-A。這不是正式 lineageId。"],
    ["完成條件", "所有學生列都有確認群組；所有班級列都有延續判定；不確定的列可保留，系統會禁止匯入並列出待處理項目。"],
    ["安全性", "這份表只產生人工確認結果，不會連線、不會寫入、不會刪除 Firebase 資料。"],
]
for row in notes:
    ws.append(row)
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 100
ws.freeze_panes = "A2"

header_fill = PatternFill("solid", fgColor="2563EB")
header_font = Font(color="FFFFFF", bold=True)
input_fill = PatternFill("solid", fgColor="FFF7CC")
for sheet in [ws]:
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

classes = wb.create_sheet("班級確認")
class_headers = ["舊班級名稱", "年級", "階段", "學校（請填）", "學年（請填）", "學期（請填）", "延續判定", "延續群組（請填）", "備註"]
classes.append(class_headers)
class_validation = DataValidation(type="list", formula1='"建立新班級線,延續既有班級線,待確認"', allow_blank=False)
classes.add_data_validation(class_validation)
for item in data["classInstances"]:
    classes.append([item["name"], item["grade"], item["stage"], item.get("school", ""), item.get("academicYear", ""), item.get("term", ""), "待確認", "", ""])
for cell in classes[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for row in range(2, classes.max_row + 1):
    class_validation.add(classes.cell(row, 7))
    for col in [4, 5, 6, 8, 9]:
        classes.cell(row, col).fill = input_fill
classes.freeze_panes = "A2"
classes.auto_filter.ref = classes.dimensions
for col, width in {"A":22,"B":12,"C":16,"D":18,"E":14,"F":12,"G":20,"H":20,"I":34}.items():
    classes.column_dimensions[col].width = width

students = wb.create_sheet("學生確認")
student_headers = ["Occurrence 顯示鍵", "舊 ID（僅供參考）", "姓名", "學校", "年級", "階段", "分組", "班級", "確認群組（請填）", "判定", "備註"]
students.append(student_headers)
judgment_validation = DataValidation(type="list", formula1='"建立新學生,與其他列相同,待確認"', allow_blank=False)
students.add_data_validation(judgment_validation)
for item in data["occurrences"]:
    students.append([item["occurrenceKey"], item["legacyStudentId"], item["name"], item["school"], item["grade"], item["stage"], item["group"], item["legacyClassId"], "", "待確認", ""])
for cell in students[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for row in range(2, students.max_row + 1):
    judgment_validation.add(students.cell(row, 10))
    students.cell(row, 9).fill = input_fill
    students.cell(row, 10).fill = input_fill
    students.cell(row, 11).fill = input_fill
students.freeze_panes = "A2"
students.auto_filter.ref = students.dimensions
for col, width in {"A":48,"B":16,"C":16,"D":14,"E":12,"F":16,"G":12,"H":30,"I":20,"J":20,"K":34}.items():
    students.column_dimensions[col].width = width
students.column_dimensions["A"].hidden = True
students.column_dimensions["B"].hidden = True

for sheet in [classes, students]:
    sheet.row_dimensions[1].height = 32
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

wb.save(output_path)
print(output_path)
